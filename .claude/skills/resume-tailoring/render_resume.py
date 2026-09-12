"""Render a resume .md into true resume typography (PDF + DOCX).

The markdown is a CONTENT spec, not a layout spec. Layout rules applied here:
  ## SECTION            -> bold heading + full-width rule beneath
  left — **right**      -> left text, right text RIGHT-ALIGNED at the margin
  - bullet              -> filled circle glyph + hanging indent
  **bold** / *italic*   -> inline runs
"""

import re
import sys

# ---------- shared parsing ----------

TOKEN = re.compile(r"(\*\*.+?\*\*|\*.+?\*)")
RIGHT = re.compile(r"^(.*?)\s+—\s+(\*\*.+?\*\*)\s*$")
# fallback: line ends in a bold date with no ' — ' marker (legacy resumes)
TRAILING_BOLD = re.compile(r"^(.*?)\s*(\*\*[^*]+\*\*)\s*$")
YEAR = re.compile(r"\b(19|20)\d{2}\b")


def parse_inline(text):
    """-> [(text, bold, italic)]"""
    runs = []
    for tok in TOKEN.split(text):
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**"):
            runs.append((tok[2:-2], True, False))
        elif tok.startswith("*") and tok.endswith("*"):
            runs.append((tok[1:-1], False, True))
        else:
            runs.append((tok, False, False))
    return runs


def split_right(text):
    """Detect 'left — **right**' -> (left_runs, right_runs) else (runs, None).

    Falls back to a trailing bold date (no em dash) so legacy resumes still get
    their dates flush right instead of dangling mid-line.
    """
    m = RIGHT.match(text)
    if m and len(m.group(2)) <= 44:
        return parse_inline(m.group(1)), parse_inline(m.group(2))

    m = TRAILING_BOLD.match(text)
    if m and m.group(1).strip():
        inner = m.group(2)[2:-2]
        if len(inner) <= 30 and YEAR.search(inner):
            return parse_inline(m.group(1).rstrip()), parse_inline(m.group(2))

    return parse_inline(text), None


def parse_md(path):
    blocks = []
    seen_name = False
    seen_contact = False
    for raw in open(path, encoding="utf-8").read().split("\n"):
        line = raw.rstrip()
        if not line or line == "---":
            continue
        if line.startswith("# "):
            blocks.append(("name", parse_inline(line[2:].strip())))
            seen_name = True
            continue
        if line.startswith("## "):
            blocks.append(("section", line[3:].strip().upper()))
            continue
        # ### / #### -> entry heading (job title, project). Never emit raw '#'.
        m = re.match(r"^#{3,6}\s+(.*)$", line)
        if m:
            head = m.group(1).strip()
            if "|" in head:
                # "Name | role | descriptor" -> bold only the name, as the
                # canonical resume format does
                first, rest = head.split("|", 1)
                head = "**" + first.strip() + "** | " + rest.strip()
            elif not head.startswith("**"):
                head = "**" + head + "**"
            left, right = split_right(head)
            blocks.append(("line", left, right))
            continue
        if line.startswith("- "):
            left, right = split_right(line[2:].strip())
            blocks.append(("bullet", left, right))
            continue
        if seen_name and not seen_contact:
            blocks.append(("contact", parse_inline(line)))
            seen_contact = True
            continue
        left, right = split_right(line)
        blocks.append(("line", left, right))
    return blocks


# ---------- PDF ----------

NAME_SIZE = 19
BODY_SIZE = 10
CONTACT_SIZE = 9.5
SECTION_SIZE = 10


def font_for(bold, italic):
    if bold and italic:
        return "Times-BoldItalic"
    if bold:
        return "Times-Bold"
    if italic:
        return "Times-Italic"
    return "Times-Roman"


def to_words(runs):
    out = []
    for text, b, i in runs:
        for w in text.split():
            out.append((w, b, i))
    return out


def runs_width(c, runs, size):
    return sum(c.stringWidth(t, font_for(b, i), size) for t, b, i in runs)


def render_pdf(blocks, out_path):
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas

    W, H = letter
    LEFT = 0.6 * 72
    RIGHT = W - 0.6 * 72
    MAXW = RIGHT - LEFT

    c = canvas.Canvas(out_path, pagesize=letter)
    y = H - 0.5 * 72

    def draw_words(words, x, y, maxw, size, leading, hang=0):
        """Greedy wrap. Returns y after the last line."""
        space = c.stringWidth(" ", "Times-Roman", size)
        line, lw = [], 0.0
        first = True
        for w, b, i in words:
            ww = c.stringWidth(w, font_for(b, i), size)
            avail = maxw if first else maxw - hang
            if line and lw + space + ww > avail:
                cx = x if first else x + hang
                for t, bb, ii in line:
                    c.setFont(font_for(bb, ii), size)
                    c.drawString(cx, y, t)
                    cx += c.stringWidth(t, font_for(bb, ii), size) + space
                y -= leading
                line, lw, first = [], 0.0, False
            line.append((w, b, i))
            lw += (space if lw else 0) + ww
        if line:
            cx = x if first else x + hang
            for t, bb, ii in line:
                c.setFont(font_for(bb, ii), size)
                c.drawString(cx, y, t)
                cx += c.stringWidth(t, font_for(bb, ii), size) + space
            y -= leading
        return y

    prev = None
    for blk in blocks:
        kind = blk[0]

        # new entry starting after a bullet list -> breathing room
        if kind == "line" and prev == "bullet":
            y -= 6
        prev = kind

        if kind == "name":
            runs = blk[1]
            wpx = runs_width(c, runs, NAME_SIZE)
            cx = (W - wpx) / 2
            for t, b, i in runs:
                c.setFont(font_for(b, i), NAME_SIZE)
                c.drawString(cx, y, t)
                cx += c.stringWidth(t, font_for(b, i), NAME_SIZE)
            y -= NAME_SIZE + 8

        elif kind == "contact":
            runs = blk[1]
            wpx = runs_width(c, runs, CONTACT_SIZE)
            cx = (W - wpx) / 2
            for t, b, i in runs:
                c.setFont(font_for(b, i), CONTACT_SIZE)
                c.drawString(cx, y, t)
                cx += c.stringWidth(t, font_for(b, i), CONTACT_SIZE)
            y -= CONTACT_SIZE + 13

        elif kind == "section":
            y -= 9
            c.setFont("Times-Bold", SECTION_SIZE)
            c.drawString(LEFT, y, blk[1])
            y -= 3
            c.setLineWidth(0.6)
            c.setStrokeGray(0.35)
            c.line(LEFT, y, RIGHT, y)
            y -= 12

        elif kind == "line":
            left, right = blk[1], blk[2]
            if right:
                rw = runs_width(c, right, BODY_SIZE)
                cx = RIGHT - rw
                for t, b, i in right:
                    c.setFont(font_for(b, i), BODY_SIZE)
                    c.drawString(cx, y, t)
                    cx += c.stringWidth(t, font_for(b, i), BODY_SIZE)
                avail = MAXW - rw - 12
            else:
                avail = MAXW
            y = draw_words(to_words(left), LEFT, y, avail, BODY_SIZE, BODY_SIZE + 2.2)

        elif kind == "bullet":
            left, right = blk[1], blk[2]
            bx = LEFT + 8
            tx = LEFT + 18
            c.setFillGray(0)
            c.circle(bx, y + 3, 1.7, stroke=0, fill=1)
            if right:
                rw = runs_width(c, right, BODY_SIZE)
                cx = RIGHT - rw
                for t, b, i in right:
                    c.setFont(font_for(b, i), BODY_SIZE)
                    c.drawString(cx, y, t)
                    cx += c.stringWidth(t, font_for(b, i), BODY_SIZE)
                avail = RIGHT - tx - rw - 12
            else:
                avail = RIGHT - tx
            y = draw_words(to_words(left), tx, y, avail, BODY_SIZE, BODY_SIZE + 2.2)

    c.save()
    print("wrote", out_path)

    # One page is non-negotiable, and this canvas does not paginate: anything
    # below the bottom margin is drawn off the sheet and silently lost. Report
    # the slack so the caller can trim (or use) it.
    bottom = 0.5 * 72
    slack = y - bottom
    lines = abs(slack) / (BODY_SIZE + 2.2)
    if slack < 0:
        print("OVERFLOW: exceeds one page by %.0fpt (~%.0f lines) - trim content"
              % (-slack, lines))
        return False
    print("fits one page: %.0fpt (~%.0f lines) of space left" % (slack, lines))
    return True


# ---------- DOCX ----------


def render_docx(blocks, out_path):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    FONT = "Times New Roman"
    doc = Document()
    for s in doc.sections:
        s.top_margin = Inches(0.5)
        s.bottom_margin = Inches(0.5)
        s.left_margin = Inches(0.6)
        s.right_margin = Inches(0.6)
    RIGHT_TAB = Inches(7.3)

    st = doc.styles["Normal"]
    st.font.name = FONT
    st.font.size = Pt(BODY_SIZE)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)

    def tight(p, before=0, after=1.5):
        pf = p.paragraph_format
        pf.space_before = Pt(before)
        pf.space_after = Pt(after)
        pf.line_spacing = 1.0

    def add(p, runs, size=BODY_SIZE):
        for t, b, i in runs:
            r = p.add_run(t)
            r.font.name = FONT
            r.font.size = Pt(size)
            r.bold = b
            r.italic = i

    def rule(p):
        pPr = p._p.get_or_add_pPr()
        bdr = OxmlElement("w:pBdr")
        bot = OxmlElement("w:bottom")
        bot.set(qn("w:val"), "single")
        bot.set(qn("w:sz"), "6")
        bot.set(qn("w:space"), "1")
        bot.set(qn("w:color"), "595959")
        bdr.append(bot)
        pPr.append(bdr)

    prev = None
    for blk in blocks:
        kind = blk[0]
        entry_gap = 5 if (kind == "line" and prev == "bullet") else 0
        prev = kind

        if kind == "name":
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add(p, blk[1], NAME_SIZE)
            tight(p, 0, 3)

        elif kind == "contact":
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add(p, blk[1], CONTACT_SIZE)
            tight(p, 0, 9)

        elif kind == "section":
            p = doc.add_paragraph()
            r = p.add_run(blk[1])
            r.font.name = FONT
            r.font.size = Pt(SECTION_SIZE)
            r.bold = True
            r.font.color.rgb = RGBColor(0, 0, 0)
            tight(p, 9, 3)
            rule(p)

        elif kind in ("line", "bullet"):
            left, right = blk[1], blk[2]
            p = doc.add_paragraph()
            pf = p.paragraph_format
            if kind == "bullet":
                pf.left_indent = Inches(0.25)
                pf.first_line_indent = Inches(-0.14)
                p.add_run("●  ").font.size = Pt(6.5)
            if right:
                pf.tab_stops.add_tab_stop(RIGHT_TAB, WD_TAB_ALIGNMENT.RIGHT)
            add(p, left)
            if right:
                p.add_run("\t")
                add(p, right)
            tight(p, entry_gap)

    doc.save(out_path)
    print("wrote", out_path)


# ---------- LETTER MODE ----------
#
# A cover letter is prose, not entries: every source line is its own paragraph
# and a blank line adds vertical space. Same fonts, margins, and centered
# name/contact header as the resume, so the two documents look like a set.


def parse_letter(path):
    blocks = []
    seen_name = seen_contact = False
    for raw in open(path, encoding="utf-8").read().split("\n"):
        line = raw.rstrip()
        if not line:
            if blocks and blocks[-1][0] != "gap":
                blocks.append(("gap",))
            continue
        if line == "---":
            continue
        if line.startswith("# ") and not seen_name:
            blocks.append(("name", parse_inline(line[2:].strip())))
            seen_name = True
            continue
        if seen_name and not seen_contact:
            blocks.append(("contact", parse_inline(line)))
            seen_contact = True
            continue
        blocks.append(("para", parse_inline(line.lstrip("> "))))
    return blocks


def render_letter_pdf(blocks, out_path):
    from reportlab.lib.pagesizes import letter as LETTER
    from reportlab.pdfgen import canvas

    W, H = LETTER
    LEFT = 0.6 * 72
    RIGHT = W - 0.6 * 72
    MAXW = RIGHT - LEFT
    LEADING = BODY_SIZE + 3.2

    c = canvas.Canvas(out_path, pagesize=LETTER)
    y = H - 0.5 * 72

    for blk in blocks:
        kind = blk[0]
        if kind == "gap":
            y -= 7
        elif kind in ("name", "contact"):
            size = NAME_SIZE if kind == "name" else CONTACT_SIZE
            runs = blk[1]
            cx = (W - runs_width(c, runs, size)) / 2
            for t, b, i in runs:
                c.setFont(font_for(b, i), size)
                c.drawString(cx, y, t)
                cx += c.stringWidth(t, font_for(b, i), size)
            y -= size + (6 if kind == "name" else 14)
        else:
            words = to_words(blk[1])
            space = c.stringWidth(" ", "Times-Roman", BODY_SIZE)
            line, lw = [], 0.0
            for w, b, i in words:
                ww = c.stringWidth(w, font_for(b, i), BODY_SIZE)
                if line and lw + space + ww > MAXW:
                    cx = LEFT
                    for t, bb, ii in line:
                        c.setFont(font_for(bb, ii), BODY_SIZE)
                        c.drawString(cx, y, t)
                        cx += c.stringWidth(t, font_for(bb, ii), BODY_SIZE) + space
                    y -= LEADING
                    line, lw = [], 0.0
                line.append((w, b, i))
                lw += (space if lw else 0) + ww
            if line:
                cx = LEFT
                for t, bb, ii in line:
                    c.setFont(font_for(bb, ii), BODY_SIZE)
                    c.drawString(cx, y, t)
                    cx += c.stringWidth(t, font_for(bb, ii), BODY_SIZE) + space
                y -= LEADING

    c.save()
    print("wrote", out_path)

    bottom = 0.5 * 72
    slack = y - bottom
    lines = abs(slack) / LEADING
    if slack < 0:
        print("OVERFLOW: exceeds one page by %.0fpt (~%.0f lines) - trim content"
              % (-slack, lines))
        return False
    print("fits one page: %.0fpt (~%.0f lines) of space left" % (slack, lines))
    return True


def render_letter_docx(blocks, out_path):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt

    FONT = "Times New Roman"
    doc = Document()
    for s in doc.sections:
        s.top_margin = Inches(0.5)
        s.bottom_margin = Inches(0.5)
        s.left_margin = Inches(0.6)
        s.right_margin = Inches(0.6)

    st = doc.styles["Normal"]
    st.font.name = FONT
    st.font.size = Pt(BODY_SIZE)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)

    pending_gap = 0
    for blk in blocks:
        kind = blk[0]
        if kind == "gap":
            pending_gap = 6
            continue
        p = doc.add_paragraph()
        size = BODY_SIZE
        if kind == "name":
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            size = NAME_SIZE
        elif kind == "contact":
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            size = CONTACT_SIZE
        for t, b, i in blk[1]:
            r = p.add_run(t)
            r.font.name = FONT
            r.font.size = Pt(size)
            r.bold = b
            r.italic = i
        pf = p.paragraph_format
        pf.space_before = Pt(pending_gap)
        pf.space_after = Pt(2 if kind == "para" else 4)
        pf.line_spacing = 1.08
        pending_gap = 0

    doc.save(out_path)
    print("wrote", out_path)


if __name__ == "__main__":
    args = sys.argv[1:]
    letter_mode = "--letter" in args
    if letter_mode:
        args.remove("--letter")

    fits = True
    if letter_mode:
        blocks = parse_letter(args[0])
        if len(args) > 1:
            fits = render_letter_pdf(blocks, args[1])
        if len(args) > 2:
            render_letter_docx(blocks, args[2])
    else:
        blocks = parse_md(args[0])
        if len(args) > 1:
            fits = render_pdf(blocks, args[1])
        if len(args) > 2:
            render_docx(blocks, args[2])
    sys.exit(0 if fits else 1)
